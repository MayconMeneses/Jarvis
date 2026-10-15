import io
import json
import types

import pytest

from jarvis import documents, files, memory
from jarvis.config import ClaudeConfig
from jarvis.dev import runs

MD = """# Relatório de vendas

Resumo do mês com **alta de 12%** e ação de promoção.

## Destaques
- Produto A cresceu
- **Produto B** caiu

1. Rever preços
2. Treinar equipe

| Mês | Total |
|---|---|
| Setembro | 1.000 |
| Outubro | 1.120 |
"""


def test_markdown_parser_recognizes_blocks():
    kinds = [b[0] for b in documents.parse_markdown(MD)]
    assert kinds == ["h", "p", "h", "ul", "ol", "table"]
    table = documents.parse_markdown(MD)[-1][1]
    assert table == [["Mês", "Total"], ["Setembro", "1.000"], ["Outubro", "1.120"]]


def test_word_document_has_headings_bold_lists_and_table(tmp_path):
    import docx

    path = documents.create("word", "relatório", MD, "Vendas 2026", folder=tmp_path)
    doc = docx.Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert path.suffix == ".docx" and "Vendas 2026" in texts and "Destaques" in texts
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "alta de 12%" in bold and "Produto B" in bold
    assert doc.tables[0].cell(2, 0).text == "Outubro"


def test_pdf_keeps_accents_and_is_readable(tmp_path):
    from pypdf import PdfReader

    path = documents.create("pdf", "relatório", MD, "Relatório mensal", folder=tmp_path)
    text = PdfReader(str(path)).pages[0].extract_text()
    assert "Relatório mensal" in text and "Mês" in text and "Treinar equipe" in text


def test_excel_keeps_numbers_formulas_and_styles_header(tmp_path):
    import openpyxl

    content = json.dumps([{"nome": "Vendas", "linhas": [["Mês", "Total"], ["Set", "1000"], ["Out", "1120,5"], ["Soma", "=SUM(B2:B3)"]]}])
    path = documents.create("excel", "vendas", content, folder=tmp_path)
    ws = openpyxl.load_workbook(str(path))["Vendas"]
    assert ws["B2"].value == 1000 and ws["B3"].value == 1120.5 and ws["B4"].value == "=SUM(B2:B3)"
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"


def test_powerpoint_has_cover_and_slides_with_notes(tmp_path):
    from pptx import Presentation

    content = json.dumps([{"titulo": "Problema", "topicos": ["Lento", "Caro"], "notas": "falar devagar"},
                          {"titulo": "Solução", "topicos": ["Jarvis"]}])
    path = documents.create("powerpoint", "apresentação", content, "Super Jarvis", folder=tmp_path)
    deck = Presentation(str(path))
    titles = [s.shapes.title.text for s in deck.slides]
    assert titles == ["Super Jarvis", "Problema", "Solução"]
    assert deck.slides[1].notes_slide.notes_text_frame.text == "falar devagar"


def test_documents_never_overwrite_and_validate(tmp_path):
    a = documents.create("word", "plano", "# A", folder=tmp_path)
    b = documents.create("word", "plano", "# B", folder=tmp_path)
    assert a.name == "plano.docx" and b.name == "plano (2).docx"
    with pytest.raises(ValueError, match="tipo inválido"):
        documents.create("exe", "x", "y", folder=tmp_path)
    with pytest.raises(ValueError, match="JSON"):
        documents.create("excel", "x", "não é json", folder=tmp_path)
    with pytest.raises(ValueError):
        files.save_output_bytes("virus.exe", b"MZ", tmp_path)


# ── pesquisa em níveis ───────────────────────────────────────────────────

def test_research_levels_and_validation(tmp_path, monkeypatch):
    monkeypatch.setattr(runs, "data_dir", lambda: tmp_path)
    store = runs.RunStore(tmp_path / "runs")
    run = runs.enqueue_research(store, "melhor framework web", "difícil")
    assert run.modo == "pesquisar" and run.nivel == "dificil" and run.projeto == "Pesquisa difícil"
    with pytest.raises(ValueError):
        runs.enqueue_research(store, "x", "infinito")
    with pytest.raises(ValueError):
        runs.enqueue_research(store, "  ", "media")
    with pytest.raises(ValueError):
        runs.enqueue(store, None, "x", "pesquisar", "y")


def test_research_runs_isolated_with_only_web_tools():
    cmd = runs.command(runs.Run("Pesquisa média", ".", "pesquisar", "tema", nivel="media"), ClaudeConfig(executable="x"))
    assert cmd[cmd.index("--tools") + 1] == "WebSearch,WebFetch"
    assert cmd[cmd.index("--allowedTools") + 1] == "WebSearch,WebFetch"
    assert "--mcp-config" not in cmd and "--strict-mcp-config" in cmd


def test_research_prompt_counts_sources_and_saves_to_obsidian(tmp_path):
    store = runs.RunStore(tmp_path / "runs")
    vault = memory.Vault(tmp_path / "cofre")
    run = store.save(runs.Run("Pesquisa média", str(tmp_path), "pesquisar", "qual banco de dados usar", nivel="media"))
    sent = {}

    def runner(cmd, **kw):
        sent.update(kw)
        result = "**Resposta**\n" + "\n".join(f"{i}. Site https://site{i}.com/artigo" for i in range(1, 13))
        return types.SimpleNamespace(returncode=0, stdout=json.dumps({"result": result}), stderr="")

    done = runs.execute(run, store, ClaudeConfig(executable="x"), runner=runner, vault=vault)
    assert "NO MÍNIMO 15 sites" in sent["input"] and "qual banco de dados usar" in sent["input"]
    assert sent["timeout"] == runs.LEVELS["media"]["timeout"]
    assert done.alteracoes.startswith("12 sites citados (meta: 15)") and "abaixo da meta" in done.alteracoes
    [report] = (vault.home / "Pesquisas").glob("*.md")
    assert "qual banco de dados usar" in report.read_text(encoding="utf-8")


def test_count_sources_ignores_duplicates_and_www():
    text = "https://www.a.com/x https://a.com/x/ https://b.org/y (https://c.dev)"
    assert runs.count_sources(text) == 3


def test_main_session_can_search_the_web(tmp_path):
    from jarvis.brains.claude_session import ClaudeSession

    cmd = ClaudeSession(ClaudeConfig(executable="x"), mcp_config=tmp_path / "m.json").command("s")
    assert cmd[cmd.index("--tools") + 1] == "WebSearch,WebFetch"
    assert cmd[cmd.index("--allowedTools") + 1] == "WebSearch,WebFetch,mcp__jarvis"
    off = ClaudeSession(ClaudeConfig(executable="x", web=False)).command("s")
    assert off[off.index("--tools") + 1] == "" and "--allowedTools" not in off


def test_title_is_not_repeated_when_markdown_starts_with_it(tmp_path):
    import docx

    path = documents.create("word", "plano", "# Plano de Estudos\n\nTexto.", "Plano de Estudos", folder=tmp_path)
    texts = [p.text for p in docx.Document(str(path)).paragraphs]
    assert texts.count("Plano de Estudos") == 1
