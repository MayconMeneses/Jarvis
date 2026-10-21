"""Documentos criados pelo celular (movel/): o JavaScript gera .docx/.xlsx/.pptx/.pdf e leitores INDEPENDENTES do Python os abrem.

Se python-docx, openpyxl, python-pptx ou pypdf não abrirem o arquivo, ele não serve para o usuário (Word, Excel, PowerPoint, Arquivos do iOS).
"""

import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest

MOVEL = Path(__file__).resolve().parent.parent / "movel"
NODE = shutil.which("node")
pytestmark = pytest.mark.skipif(NODE is None, reason="node não está instalado")

MARKDOWN = (
    "# Relatório de Vendas\nTexto com **negrito**, ação e “aspas” & <tags>.\n\n## Lista\n- um\n- dois\n1. primeiro\n2. segundo\n\n"
    "## Dados\n| Produto | Qtd | Código |\n|---|---|---|\n| Maçã | 3 | 007 |\n| Pão | 5,5 | =1+1 |\n\n## Outra seção\nFim."
)
JS = """
import { readFileSync, writeFileSync } from "node:fs";
import { criarArquivo } from "./js/doc-proposta.js";
const saida = process.argv[2], md = readFileSync(`${saida}/entrada.md`, "utf8");
for (const tipo of ["docx", "xlsx", "pptx", "pdf"]) {
  const a = criarArquivo(tipo, md, new Date(2026, 9, 21, 10, 0, 0));
  writeFileSync(`${saida}/doc.${tipo}`, a.bytes);
}
"""


@pytest.fixture(scope="module")
def arquivos(tmp_path_factory):
    pasta = tmp_path_factory.mktemp("docs")
    script = pasta / "gerar.mjs"
    # o script fica na pasta temporária (não dentro de movel/); os imports usam o caminho absoluto de movel/js
    script.write_text(JS.replace('"./js/', f'"{(MOVEL / "js").as_uri()}/'), encoding="utf-8")
    (pasta / "entrada.md").write_text(MARKDOWN, encoding="utf-8")  # arquivo, não argumento: o Windows estraga acentos na linha de comando
    saida = subprocess.run([NODE, str(script), str(pasta)], capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert saida.returncode == 0, saida.stderr
    return pasta


@pytest.mark.parametrize("tipo", ["docx", "xlsx", "pptx"])
def test_zip_is_valid_and_content_types_come_first(arquivos, tipo):
    with zipfile.ZipFile(arquivos / f"doc.{tipo}") as z:
        assert z.testzip() is None
        assert z.namelist()[0] == "[Content_Types].xml"


def test_python_docx_opens_the_word_file(arquivos):
    import docx

    d = docx.Document(str(arquivos / "doc.docx"))
    estilos = [(p.style.name, p.text) for p in d.paragraphs if p.text]
    assert estilos[0] == ("Heading 1", "Relatório de Vendas")
    assert "negrito, ação e “aspas” & <tags>." in estilos[1][1]
    assert [t for e, t in estilos if e == "Heading 2"] == ["Lista", "Dados", "Outra seção"]
    assert [[c.text for c in linha.cells] for linha in d.tables[0].rows] == [["Produto", "Qtd", "Código"], ["Maçã", "3", "007"], ["Pão", "5,5", "=1+1"]]
    assert d.tables[0].rows[0].cells[0].paragraphs[0].runs[0].bold
    assert d.core_properties.title == "Relatório de Vendas"
    assert any(r.bold for p in d.paragraphs for r in p.runs if r.text == "negrito")


def test_openpyxl_opens_the_excel_file_and_text_never_becomes_formula(arquivos):
    import openpyxl

    wb = openpyxl.load_workbook(arquivos / "doc.xlsx")
    assert wb.sheetnames == ["Dados"]
    ws = wb["Dados"]
    assert [c.value for c in ws[1]] == ["Produto", "Qtd", "Código"]
    assert ws["B2"].value == 3 and ws["B2"].data_type == "n"
    assert ws["B3"].value == 5.5
    assert ws["C2"].value == "007" and ws["C2"].data_type == "s"
    assert ws["C3"].value == "=1+1" and ws["C3"].data_type == "s"  # texto, não fórmula
    assert ws["A1"].font.b is True
    assert ws.freeze_panes == "A2"


def test_python_pptx_opens_the_powerpoint_file(arquivos):
    from pptx import Presentation

    p = Presentation(str(arquivos / "doc.pptx"))
    titulos = [s.shapes.title.text for s in p.slides]
    assert titulos == ["Relatório de Vendas", "Lista", "Dados", "Outra seção"]
    corpo = [sh.text_frame.text for sh in p.slides[1].placeholders if sh.placeholder_format.idx == 1]
    assert corpo == ["um\ndois\nprimeiro\nsegundo"]
    assert "Maçã | 3 | 007" in p.slides[2].placeholders[1].text_frame.text
    assert p.slide_width == 12192000 and p.slide_height == 6858000


def test_pypdf_extracts_text_from_the_pdf(arquivos):
    from pypdf import PdfReader

    r = PdfReader(str(arquivos / "doc.pdf"))
    assert len(r.pages) == 1
    texto = r.pages[0].extract_text()
    assert "Relatório de Vendas" in texto
    assert "ação" in texto and "“aspas”" in texto and "& <tags>" in texto
    assert "Maçã" in texto and "007" in texto
    assert r.metadata.title == "Relatório de Vendas"
