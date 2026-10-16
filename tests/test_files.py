import io
import os
import time

import pytest

from jarvis import files

def mini_pdf(text: str) -> bytes:
    """PDF mínimo válido (com xref) contendo uma linha de texto."""
    stream = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 100]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return bytes(out)


MINI_PDF = mini_pdf("Relatorio de vendas")


@pytest.fixture
def store(tmp_path):
    return files.AttachmentStore(tmp_path / "anexos")


def test_text_and_code_are_read_with_encoding_fallback(store):
    att = store.add("script.py", "print('olá')\n".encode("cp1252"))
    assert att.tipo == "texto" and "print('olá')" in files.read_text(att)


def test_pdf_text_is_extracted(store):
    att = store.add("relatorio.pdf", MINI_PDF)
    assert att.tipo == "pdf" and "Relatorio de vendas" in files.read_text(att)


def test_word_paragraphs_and_tables_are_extracted(store):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Proposta comercial")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Plano", "R$ 99"
    buffer = io.BytesIO()
    doc.save(buffer)
    text = files.read_text(store.add("proposta.docx", buffer.getvalue()))
    assert "Proposta comercial" in text and "Plano | R$ 99" in text


def test_excel_sheets_are_extracted(store):
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Vendas"
    book.active.append(["Mês", "Total"])
    book.active.append(["Outubro", 1234])
    buffer = io.BytesIO()
    book.save(buffer)
    text = files.read_text(store.add("vendas.xlsx", buffer.getvalue()))
    assert "--- planilha Vendas ---" in text and "Outubro | 1234" in text


def test_images_are_resized_png(store):
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (4000, 1000), "red").save(buffer, format="JPEG")
    att = store.add("tela.jpg", buffer.getvalue())
    png = files.image_bytes(att)
    with Image.open(io.BytesIO(png)) as img:
        assert img.format == "PNG" and max(img.size) == 1568


def test_long_text_is_cut_with_notice(store):
    att = store.add("grande.txt", b"a" * 100)
    assert "cortado" in files.read_text(att, max_chars=10)


@pytest.mark.parametrize("name,data,error", [
    ("programa.exe", b"MZ", "não suportado"),
    ("vazio.txt", b"", "vazio"),
])
def test_rejects_unsupported_and_empty(store, name, data, error):
    with pytest.raises(ValueError, match=error):
        store.add(name, data)


def test_size_limit(store, monkeypatch):
    monkeypatch.setattr(files, "MAX_BYTES", 10)
    with pytest.raises(ValueError, match="limite"):
        store.add("a.txt", b"x" * 11)


def test_names_are_sanitized_and_get_validates_ids(store):
    att = store.add("..\\..\\segredo.txt", b"x")
    assert att.nome == "segredo.txt" and store.folder in __import__("pathlib").Path(att.caminho).parents
    assert store.get(att.id).nome == "segredo.txt"
    for bad in ("../x", "abc", "zzzzzzzz"):
        with pytest.raises(ValueError):
            store.get(bad)


def test_old_attachments_are_cleaned(store):
    old = store.add("velho.txt", b"x")
    past = time.time() - 31 * 86400
    os.utime(old.caminho, (past, past))
    store.add("novo.txt", b"y")
    assert not os.path.exists(old.caminho)


def test_save_output_never_overwrites_and_only_text(tmp_path):
    first = files.save_output("versao melhorada.py", "x = 1", tmp_path)
    second = files.save_output("versao melhorada.py", "x = 2", tmp_path)
    assert first.read_text(encoding="utf-8") == "x = 1" and second.name == "versao melhorada (2).py"
    assert files.save_output("notas", "oi", tmp_path).suffix == ".md"
    with pytest.raises(ValueError):
        files.save_output("virus.exe", "MZ", tmp_path)
    assert files.save_output("..\\..\\fora.txt", "x", tmp_path).parent == tmp_path


def test_message_lists_attachment_ids_for_the_ai(tmp_path):
    pytest.importorskip("webview")
    from jarvis.brains.base import Answer
    from jarvis.brains.router import Router
    from jarvis.chat import Conversation
    from jarvis.config import Config
    from jarvis.ui.api import Api

    seen = []

    class Brain:
        name = "claude"

        def ask(self, messages, system):
            seen.append(messages[-1]["content"])
            return Answer("ok", "claude")

    api = Api(Config(), Conversation(Router([Brain()]), "sys"))
    api._files = files.AttachmentStore(tmp_path)
    att = api._files.add("dados.csv", b"a,b\n1,2")
    api.send("", [att.id])
    assert "Analise os arquivos anexados." in seen[0]
    assert f"dados.csv (texto, 1 KB) — id {att.id}" in seen[0] and "ler_anexo" in seen[0]


def test_mcp_tool_returns_image_for_pictures(tmp_path, monkeypatch):
    from PIL import Image as PILImage

    from jarvis.mcp_tools import document_tools as mcp_server

    store = files.AttachmentStore(tmp_path)
    monkeypatch.setattr(files, "AttachmentStore", lambda: store)
    buffer = io.BytesIO()
    PILImage.new("RGB", (10, 10)).save(buffer, format="PNG")
    pic = store.add("erro.png", buffer.getvalue())
    txt = store.add("nota.md", b"# Oi")
    assert isinstance(mcp_server.ler_anexo(pic.id), mcp_server.Image)
    assert mcp_server.ler_anexo(txt.id).endswith("# Oi")
    assert "Não foi possível" in mcp_server.ler_anexo("00000000")
