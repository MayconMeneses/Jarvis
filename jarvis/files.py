"""Anexos para a Jarvis ler, e arquivos que ela cria.

- Anexos ficam em %LOCALAPPDATA%\\SuperJarvis\\anexos (cópias; o original não é tocado)
  e são apagados depois de 30 dias.
- A IA lê anexos só pelo id (ferramenta ler_anexo): nunca recebe caminhos livres do PC.
- Arquivos criados pela Jarvis vão para Documentos\\Jarvis, sempre com nome novo
  (nunca sobrescreve nada).
"""

from __future__ import annotations

import io
import re
import time
import uuid
from dataclasses import asdict, dataclass
from pathlib import Path

from jarvis.config import data_dir

MAX_BYTES = 25 * 1024 * 1024
MAX_CHARS = 60_000
KEEP_DAYS = 30
TEXT_EXT = {
    ".txt", ".md", ".py", ".js", ".ts", ".tsx", ".jsx", ".json", ".csv", ".html", ".htm", ".css",
    ".sql", ".java", ".kt", ".c", ".h", ".cpp", ".cs", ".go", ".rs", ".php", ".rb", ".yml", ".yaml",
    ".toml", ".ini", ".cfg", ".xml", ".log", ".bat", ".ps1", ".sh", ".vue", ".svelte", ".dart", ".swift",
}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
DOC_EXT = {".pdf": "pdf", ".docx": "word", ".xlsx": "excel"}
_UNSAFE = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def kind_of(name: str) -> str:
    ext = Path(name).suffix.lower()
    if ext in TEXT_EXT:
        return "texto"
    if ext in IMAGE_EXT:
        return "imagem"
    if ext in DOC_EXT:
        return DOC_EXT[ext]
    raise ValueError(f"tipo de arquivo não suportado: {ext or 'sem extensão'} "
                     "(aceitos: texto/código, PDF, Word .docx, Excel .xlsx e imagens)")


def safe_name(name: str) -> str:
    name = _UNSAFE.sub("_", Path(name).name).strip(" .") or "arquivo"
    return name[:120]


@dataclass
class Attachment:
    id: str
    nome: str
    tipo: str
    tamanho_kb: int
    caminho: str

    def public(self) -> dict:
        data = asdict(self)
        del data["caminho"]
        return data


class AttachmentStore:
    def __init__(self, folder: Path | None = None) -> None:
        self.folder = folder or data_dir() / "anexos"

    def add(self, name: str, data: bytes) -> Attachment:
        name = safe_name(name)
        kind = kind_of(name)
        if len(data) > MAX_BYTES:
            raise ValueError(f"'{name}' tem {len(data) // 1024 // 1024} MB; o limite é 25 MB")
        if not data:
            raise ValueError(f"'{name}' está vazio")
        self.folder.mkdir(parents=True, exist_ok=True)
        self._cleanup()
        att_id = uuid.uuid4().hex[:8]
        path = self.folder / f"{att_id}__{name}"
        path.write_bytes(data)
        return Attachment(att_id, name, kind, max(1, len(data) // 1024), str(path))

    def add_path(self, path: str | Path) -> Attachment:
        path = Path(path)
        if not path.is_file():
            raise ValueError(f"arquivo não encontrado: {path}")
        if path.stat().st_size > MAX_BYTES:
            raise ValueError(f"'{path.name}' passa do limite de 25 MB")
        return self.add(path.name, path.read_bytes())

    def get(self, att_id: str) -> Attachment:
        if not re.fullmatch(r"[0-9a-f]{8}", att_id or ""):
            raise ValueError("id de anexo inválido")
        found = list(self.folder.glob(f"{att_id}__*")) if self.folder.is_dir() else []
        if not found:
            raise ValueError("anexo não encontrado (pode ter expirado)")
        path = found[0]
        name = path.name.split("__", 1)[1]
        return Attachment(att_id, name, kind_of(name), max(1, path.stat().st_size // 1024), str(path))

    def _cleanup(self) -> None:
        cutoff = time.time() - KEEP_DAYS * 86400
        for path in self.folder.glob("*__*"):
            try:
                if path.stat().st_mtime < cutoff:
                    path.unlink()
            except OSError:
                pass


# ── leitura ──────────────────────────────────────────────────────────────

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_text(att: Attachment, max_chars: int = MAX_CHARS) -> str:
    path = Path(att.caminho)
    if att.tipo == "texto":
        text = _decode(path.read_bytes())
    elif att.tipo == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [f"--- página {i} ---\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1)]
        text = "\n\n".join(pages)
        if not text.replace("--- página", "").strip(" -0123456789\n"):
            text += "\n\n(O PDF parece ser só imagem/escaneado: não há texto para extrair.)"
    elif att.tipo == "word":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(parts)
    elif att.tipo == "excel":
        import openpyxl

        book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for sheet in book.worksheets:
            parts.append(f"--- planilha {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True, max_row=2000):
                if any(v is not None for v in row):
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
        book.close()
        text = "\n".join(parts)
    else:
        raise ValueError("use image_bytes para imagens")
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n[... cortado: o arquivo tem {len(text)} caracteres; mostrados os primeiros {max_chars}]"
    return text


def image_bytes(att: Attachment, max_side: int = 1568) -> bytes:
    """Imagem em PNG, reduzida para o tamanho que o modelo aproveita."""
    from PIL import Image

    with Image.open(att.caminho) as img:
        img = img.convert("RGB") if img.mode not in ("RGB", "L") else img
        img.thumbnail((max_side, max_side))
        out = io.BytesIO()
        img.save(out, format="PNG", optimize=True)
        return out.getvalue()


# ── arquivos criados pela Jarvis ─────────────────────────────────────────

def output_dir() -> Path:
    return Path.home() / "Documents" / "Jarvis"


def save_output(name: str, content: str, folder: Path | None = None) -> Path:
    name = safe_name(name)
    if not Path(name).suffix:
        name += ".md"
    if Path(name).suffix.lower() not in TEXT_EXT:
        raise ValueError("a Jarvis só cria arquivos de texto/código (ex.: .md, .txt, .py, .html)")
    folder = folder or output_dir()
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / name
    counter = 2
    while path.exists():  # nunca sobrescreve
        path = folder / f"{Path(name).stem} ({counter}){Path(name).suffix}"
        counter += 1
    path.write_text(content, encoding="utf-8")
    return path
