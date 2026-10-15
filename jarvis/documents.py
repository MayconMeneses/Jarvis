"""Criação de documentos: Word (.docx), PDF, Excel (.xlsx) e PowerPoint (.pptx).

Word e PDF recebem Markdown simples (títulos #, listas -, 1., tabelas |a|b|, **negrito**).
Excel recebe planilhas e PowerPoint recebe slides, ambos em JSON.
Tudo é salvo em Documentos\\Jarvis, sem sobrescrever (ver files.save_output_bytes).
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

from jarvis import files

KINDS = {"word": ".docx", "pdf": ".pdf", "excel": ".xlsx", "powerpoint": ".pptx"}
_FONT_DIR = Path(r"C:\Windows\Fonts")


# ── Markdown simples ─────────────────────────────────────────────────────

def parse_markdown(text: str) -> list[tuple]:
    """Blocos: ("h", nível, texto) ("p", texto) ("ul", [itens]) ("ol", [itens]) ("table", [linhas])."""
    blocks: list[tuple] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(("p", " ".join(paragraph)))
            paragraph.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue
        if m := re.match(r"^(#{1,4})\s+(.*)", line):
            flush()
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
        elif m := re.match(r"^[-*•]\s+(.*)", line):
            flush()
            if blocks and blocks[-1][0] == "ul":
                blocks[-1][1].append(m.group(1))
            else:
                blocks.append(("ul", [m.group(1)]))
        elif m := re.match(r"^\d+[.)]\s+(.*)", line):
            flush()
            if blocks and blocks[-1][0] == "ol":
                blocks[-1][1].append(m.group(1))
            else:
                blocks.append(("ol", [m.group(1)]))
        elif line.startswith("|") and line.endswith("|"):
            flush()
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue  # linha separadora |---|---|
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
        else:
            paragraph.append(line)
    flush()
    return blocks


def body_blocks(title: str, markdown: str) -> list[tuple]:
    """Blocos do texto, sem repetir o título quando o Markdown começa com o mesmo # título."""
    blocks = parse_markdown(markdown)
    if title and blocks and blocks[0][0] == "h" and blocks[0][2].strip().lower() == title.strip().lower():
        blocks = blocks[1:]
    return blocks


def bold_runs(text: str) -> list[tuple[str, bool]]:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, False) for p in parts if p]


# ── Word ─────────────────────────────────────────────────────────────────

def make_docx(title: str, markdown: str) -> bytes:
    import docx

    document = docx.Document()
    if title:
        document.add_heading(title, level=0)
    for block in body_blocks(title, markdown):
        kind = block[0]
        if kind == "h":
            document.add_heading(block[2], level=min(block[1], 4))
        elif kind == "p":
            par = document.add_paragraph()
            for text, bold in bold_runs(block[1]):
                par.add_run(text).bold = bold
        elif kind in ("ul", "ol"):
            style = "List Bullet" if kind == "ul" else "List Number"
            for item in block[1]:
                par = document.add_paragraph(style=style)
                for text, bold in bold_runs(item):
                    par.add_run(text).bold = bold
        elif kind == "table":
            rows = block[1]
            width = max(len(r) for r in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j in range(width):
                    cell = table.cell(i, j)
                    cell.text = row[j] if j < len(row) else ""
                    if i == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────

def make_pdf(title: str, markdown: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    regular, bold = _FONT_DIR / "segoeui.ttf", _FONT_DIR / "segoeuib.ttf"
    if regular.exists() and bold.exists():  # acentos e símbolos Unicode
        pdf.add_font("Doc", "", str(regular))
        pdf.add_font("Doc", "B", str(bold))
        family = "Doc"
    else:
        family = "Helvetica"
    pdf.add_page()

    def write(text: str, size: float = 11, heading: bool = False) -> None:
        pdf.set_font(family, "B" if heading else "", size)
        rich = not heading and "**" in text  # fpdf2 entende **negrito** quando markdown=True
        pdf.multi_cell(0, size * 0.55, text.replace("**", "") if heading else text,
                       markdown=rich, new_x="LMARGIN", new_y="NEXT")

    if title:
        write(title, 20, heading=True)
        pdf.ln(3)
    for block in body_blocks(title, markdown):
        kind = block[0]
        if kind == "h":
            pdf.ln(2)
            write(block[2], {1: 16, 2: 14, 3: 12.5}.get(block[1], 12), heading=True)
            pdf.ln(1)
        elif kind == "p":
            write(block[1])
            pdf.ln(2)
        elif kind in ("ul", "ol"):
            for n, item in enumerate(block[1], 1):
                write(("• " if kind == "ul" else f"{n}. ") + item)
            pdf.ln(2)
        elif kind == "table":
            pdf.set_font(family, "", 10)
            with pdf.table(text_align="LEFT") as table:
                for row in block[1]:
                    table_row = table.row()
                    for cell in row:
                        table_row.cell(cell.replace("**", ""))
            pdf.ln(2)
    return bytes(pdf.output())


# ── Excel ────────────────────────────────────────────────────────────────

def make_xlsx(sheets: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    book = openpyxl.Workbook()
    book.remove(book.active)
    for n, sheet in enumerate(sheets or [{"nome": "Planilha1", "linhas": []}], 1):
        ws = book.create_sheet(str(sheet.get("nome") or f"Planilha{n}")[:31])
        for row in sheet.get("linhas", []):
            ws.append([_cell(v) for v in row])
        if ws.max_row >= 1 and sheet.get("cabecalho", True):
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            ws.freeze_panes = "A2"
        for column in ws.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            ws.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 60)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _cell(value):
    """Mantém números como números (para fórmulas e somas funcionarem no Excel)."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if text.startswith("="):
        return text  # fórmula
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+[.,]\d+", text):
        return float(text.replace(",", "."))
    return value


# ── PowerPoint ───────────────────────────────────────────────────────────

def make_pptx(title: str, slides: list[dict]) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = 12192000, 6858000  # 16:9
    if title:
        cover = deck.slides.add_slide(deck.slide_layouts[0])
        cover.shapes.title.text = title
        if len(cover.placeholders) > 1:
            cover.placeholders[1].text = ""
    for spec in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(spec.get("titulo", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, topic in enumerate(spec.get("topicos", [])):
            par = body.paragraphs[0] if i == 0 else body.add_paragraph()
            par.text = str(topic).replace("**", "")
            par.font.size = Pt(22)
        if spec.get("notas"):
            slide.notes_slide.notes_text_frame.text = str(spec["notas"])
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


# ── entrada única ────────────────────────────────────────────────────────

def create(kind: str, name: str, content: str, title: str = "", folder: Path | None = None) -> Path:
    kind = kind.strip().lower()
    if kind not in KINDS:
        raise ValueError(f"tipo inválido: {kind} (use word, pdf, excel ou powerpoint)")
    if kind in ("word", "pdf"):
        data = make_docx(title, content) if kind == "word" else make_pdf(title, content)
    else:
        try:
            parsed = json.loads(content)
        except json.JSONDecodeError as exc:
            raise ValueError(f"para {kind}, o conteúdo deve ser JSON válido ({exc.msg})") from exc
        if kind == "excel":
            sheets = parsed if isinstance(parsed, list) else parsed.get("planilhas", [parsed])
            data = make_xlsx(sheets)
        else:
            slides = parsed if isinstance(parsed, list) else parsed.get("slides", [])
            data = make_pptx(title or (parsed.get("titulo", "") if isinstance(parsed, dict) else ""), slides)
    stem = Path(files.safe_name(name)).stem or "documento"
    return files.save_output_bytes(stem + KINDS[kind], data, folder)
